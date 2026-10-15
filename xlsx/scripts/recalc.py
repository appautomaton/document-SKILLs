# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Recalculate every formula in an .xlsx or .xlsm workbook and report problems.

Usage:
    uv run recalc.py <workbook> [--timeout SECONDS] [--force]

LibreOffice recalculates the workbook on load and the result is written back
in place, so cached values are current wherever the file is opened next.
Prints a JSON report. Exit status: 0 clean, 1 problems found, 2 could not run.

Checked before recalculating, because LibreOffice would mask or worsen them:
  - Newer Excel functions written without their _xlfn. prefix. They would come
    back as #NAME? with the formula lowercased, so the file is left untouched
    and the corrected formulas are reported instead.
  - Links to other workbooks, whose values recalculation would replace.
    Refused unless --force.

Reported after recalculating: error values by type and location, circular
references (direct and named-range references; INDIRECT and OFFSET are not
followed), and formulas LibreOffice left without a value.
"""

import argparse
import json
import shutil
import sys
import tempfile
import zipfile
from bisect import bisect_left, bisect_right
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, str(Path(__file__).parent))
import soffice  # noqa: E402

MAX_LOCATIONS = 50
RECALC_ON_LOAD = {"/org.openoffice.Office.Calc/Formula/Load": {"OOXMLRecalcMode": 0}}
FILTERS = {
    ".xlsx": "xlsx:Calc MS Excel 2007 XML",
    ".xlsm": "xlsm:Calc MS Excel 2007 VBA XML",
}

# Functions added after Excel 2007. The file format stores them with an
# _xlfn. prefix (FILTER and SORT with _xlfn._xlws.); without it, neither Excel
# nor LibreOffice recognises them.
XLWS = {"FILTER", "SORT"}
XLFN = XLWS | set(
    """
    ACOT ACOTH AGGREGATE ARABIC ARRAYTOTEXT BASE BETA.DIST BETA.INV BINOM.DIST
    BINOM.DIST.RANGE BINOM.INV BITAND BITLSHIFT BITOR BITRSHIFT BITXOR BYCOL BYROW
    CEILING.MATH CEILING.PRECISE CHISQ.DIST CHISQ.DIST.RT CHISQ.INV CHISQ.INV.RT
    CHISQ.TEST CHOOSECOLS CHOOSEROWS COMBINA CONCAT CONFIDENCE.NORM CONFIDENCE.T COT
    COTH COVARIANCE.P COVARIANCE.S CSC CSCH DAYS DECIMAL DROP ENCODEURL ERF.PRECISE
    ERFC.PRECISE EXPAND EXPON.DIST F.DIST F.DIST.RT F.INV F.INV.RT F.TEST FILTERXML
    FLOOR.MATH FLOOR.PRECISE FORECAST.ETS FORECAST.ETS.CONFINT FORECAST.ETS.SEASONALITY
    FORECAST.ETS.STAT FORECAST.LINEAR FORMULATEXT GAMMA GAMMA.DIST GAMMA.INV
    GAMMALN.PRECISE GAUSS GROUPBY HSTACK HYPGEOM.DIST IFNA IFS IMAGE IMCOSH IMCOT
    IMCSC IMCSCH IMSEC IMSECH IMSINH IMTAN ISFORMULA ISOMITTED ISOWEEKNUM LAMBDA LET
    LOGNORM.DIST LOGNORM.INV MAKEARRAY MAP MAXIFS MINIFS MODE.MULT MODE.SNGL MUNIT
    NEGBINOM.DIST NETWORKDAYS.INTL NORM.DIST NORM.INV NORM.S.DIST NORM.S.INV
    NUMBERVALUE PDURATION PERCENTILE.EXC PERCENTILE.INC PERCENTOF PERCENTRANK.EXC
    PERCENTRANK.INC PERMUTATIONA PHI PIVOTBY POISSON.DIST QUARTILE.EXC QUARTILE.INC
    RANDARRAY RANK.AVG RANK.EQ REDUCE REGEXEXTRACT REGEXREPLACE REGEXTEST RRI SCAN
    SEC SECH SEQUENCE SHEET SHEETS SKEW.P SORTBY STDEV.P STDEV.S SWITCH T.DIST
    T.DIST.2T T.DIST.RT T.INV T.INV.2T T.TEST TAKE TEXTAFTER TEXTBEFORE TEXTJOIN
    TEXTSPLIT TOCOL TOROW TRIMRANGE UNICHAR UNICODE UNIQUE VALUETOTEXT VAR.P VAR.S
    VSTACK WEBSERVICE WEIBULL.DIST WORKDAY.INTL WRAPCOLS WRAPROWS XLOOKUP XMATCH XOR
    Z.TEST
    """.split()
)


def formula_text(value):
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def read_formulas(path):
    """Return {(sheet, coordinate): formula} for every formula cell."""
    wb = load_workbook(path)
    found = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if (text := formula_text(cell.value)) is not None:
                    found[(ws.title, cell.coordinate)] = text
    names = {name: dn.attr_text for name, dn in wb.defined_names.items()}
    wb.close()
    return found, names


def add_prefixes(formula):
    """Return (formula with missing prefixes added, names of functions fixed)."""
    tokenizer = Tokenizer(formula)
    fixed = []
    for token in tokenizer.items:
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper()
            if name in XLFN:
                fixed.append(name)
                prefix = "_xlfn._xlws." if name in XLWS else "_xlfn."
                token.value = prefix + token.value
    return tokenizer.render(), fixed


def missing_prefixes(formulas):
    issues = []
    for (sheet, coord), formula in formulas.items():
        try:
            fix, names = add_prefixes(formula)
        except Exception:
            continue
        if names:
            issues.append(
                {"cell": f"{sheet}!{coord}", "functions": names, "formula": formula, "fix": fix}
            )
    return issues


def split_ref(ref, sheet):
    if "!" not in ref:
        return sheet, ref
    where, ref = ref.rsplit("!", 1)
    if where.startswith("'"):
        where = where[1:-1].replace("''", "'")
    return where, ref


def referenced_ranges(formula, sheet, names, depth=0):
    """Yield (sheet, min_col, min_row, max_col, max_row) for each reference."""
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE" or token.value.startswith("["):
            continue
        target_sheet, ref = split_ref(token.value, sheet)
        try:
            yield (target_sheet, *range_boundaries(ref.replace("$", "")))
        except ValueError:
            if depth < 5 and (target := names.get(token.value)):
                yield from referenced_ranges("=" + target, sheet, names, depth + 1)


def circular_references(formulas, names):
    """Return the strongly connected groups of formula cells that reference each other."""
    index = {}  # sheet -> col -> sorted rows of formula cells
    for sheet, coord in formulas:
        col, row = coordinate_from_string(coord)
        index.setdefault(sheet, {}).setdefault(column_index_from_string(col), []).append(row)
    for cols in index.values():
        for rows in cols.values():
            rows.sort()

    graph = {}
    for (sheet, coord), formula in formulas.items():
        deps = set()
        for target, c1, r1, c2, r2 in referenced_ranges(formula, sheet, names):
            cols = index.get(target, {})
            for col in cols if c1 is None else range(c1, c2 + 1):
                rows = cols.get(col)
                if not rows:
                    continue
                lo = 0 if r1 is None else bisect_left(rows, r1)
                hi = len(rows) if r2 is None else bisect_right(rows, r2)
                deps.update((target, f"{get_column_letter(col)}{r}") for r in rows[lo:hi])
        graph[(sheet, coord)] = deps

    # Tarjan's algorithm, iterative so deep dependency chains cannot overflow.
    order, low, on_stack, stack, groups = {}, {}, set(), [], []
    for root in graph:
        if root in order:
            continue
        work = [(root, iter(graph[root]))]
        order[root] = low[root] = len(order)
        stack.append(root)
        on_stack.add(root)
        while work:
            node, children = work[-1]
            for child in children:
                if child not in order:
                    order[child] = low[child] = len(order)
                    stack.append(child)
                    on_stack.add(child)
                    work.append((child, iter(graph.get(child, ()))))
                    break
                if child in on_stack:
                    low[node] = min(low[node], order[child])
            else:
                work.pop()
                if work:
                    low[work[-1][0]] = min(low[work[-1][0]], low[node])
                if low[node] == order[node]:
                    group = []
                    while True:
                        member = stack.pop()
                        on_stack.discard(member)
                        group.append(member)
                        if member == node:
                            break
                    if len(group) > 1 or node in graph[node]:
                        groups.append(sorted(f"{s}!{c}" for s, c in group))
    return groups


def scan_values(path, formulas):
    """Return ({error: [locations]}, [uncalculated formula locations])."""
    wb = load_workbook(path, data_only=True)
    errors, uncalculated = {}, []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                where = f"{ws.title}!{cell.coordinate}"
                if cell.data_type == "e":
                    errors.setdefault(str(cell.value), []).append(where)
                elif cell.value is None and cell.data_type == "n" and (ws.title, cell.coordinate) in formulas:
                    uncalculated.append(where)
    wb.close()
    return errors, uncalculated


def summarize(locations):
    return {"count": len(locations), "locations": locations[:MAX_LOCATIONS]}


def recalc(path, timeout, force):
    if not path.is_file():
        return {"status": "failed", "error": f"{path} does not exist"}
    if path.suffix.lower() not in FILTERS:
        return {
            "status": "failed",
            "error": f"recalc handles .xlsx and .xlsm only, not {path.suffix or 'this file'}",
        }

    formulas, names = read_formulas(path)
    report = {"file": str(path), "total_formulas": len(formulas)}

    if issues := missing_prefixes(formulas):
        return report | {
            "status": "errors_found",
            "recalculated": False,
            "missing_prefix": issues[:MAX_LOCATIONS],
            "note": "Rewrite these formulas as shown in 'fix', then run recalc again. "
            "The file was not changed.",
        }

    with zipfile.ZipFile(path) as zf:
        linked = any(n.startswith("xl/externalLinks/") for n in zf.namelist())
    if linked and not force:
        return report | {
            "status": "failed",
            "recalculated": False,
            "error": "The workbook links to other workbooks. Recalculating would replace "
            "the linked values with whatever LibreOffice can resolve. "
            "Pass --force to recalculate anyway.",
        }

    with tempfile.TemporaryDirectory() as tmp:
        try:
            out = soffice.convert(
                path, tmp, FILTERS[path.suffix.lower()], timeout=timeout, settings=RECALC_ON_LOAD
            )
        except soffice.SofficeError as e:
            return report | {"status": "failed", "recalculated": False, "error": str(e)}
        shutil.copyfile(out, path)

    errors, uncalculated = scan_values(path, formulas)
    cycles = circular_references(formulas, names)
    total_errors = sum(len(v) for v in errors.values())
    report |= {
        "status": "errors_found" if total_errors or cycles or uncalculated else "success",
        "recalculated": True,
        "total_errors": total_errors,
        "errors": {kind: summarize(locs) for kind, locs in sorted(errors.items())},
    }
    if cycles:
        report["circular_references"] = cycles[:MAX_LOCATIONS]
    if uncalculated:
        report["uncalculated"] = summarize(uncalculated)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--timeout", type=int, default=120, help="seconds (default 120)")
    parser.add_argument("--force", action="store_true", help="recalculate despite external links")
    args = parser.parse_args()

    report = recalc(args.workbook.resolve(), args.timeout, args.force)
    print(json.dumps(report, indent=2, ensure_ascii=False))
    sys.exit({"success": 0, "errors_found": 1}.get(report["status"], 2))


if __name__ == "__main__":
    main()
