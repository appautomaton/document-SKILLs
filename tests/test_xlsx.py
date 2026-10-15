import subprocess
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

RECALC = "xlsx/scripts/recalc.py"


def book(path, cells, names=None):
    wb = Workbook()
    wb.active.title = "S"
    for ref, value in cells.items():
        sheet, _, coord = ref.rpartition("!")
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet) if sheet else wb.active
        ws[coord] = value
    for name, target in (names or {}).items():
        wb.defined_names[name] = DefinedName(name, attr_text=target)
    wb.save(path)
    return path


def test_clean_workbook_gets_values(run, tmp_path):
    path = book(tmp_path / "ok.xlsx", {"A1": 2, "A2": "=A1*21", "O!A1": "=S!A2+1"})
    result = run(RECALC, path)
    report = result.json()
    assert result.code == 0, result.out
    assert report["status"] == "success" and report["total_formulas"] == 2
    wb = load_workbook(path, data_only=True)
    assert wb["S"]["A2"].value == 42 and wb["O"]["A1"].value == 43


def test_errors_are_located_and_text_is_not_mistaken_for_one(run, tmp_path):
    path = book(tmp_path / "e.xlsx", {"A1": 0, "A2": "=1/A1", "A3": "note: #N/A means missing"})
    result = run(RECALC, path)
    report = result.json()
    assert result.code == 1
    assert report["errors"] == {"#DIV/0!": {"count": 1, "locations": ["S!A2"]}}


def test_missing_xlfn_prefix_is_reported_and_file_untouched(run, tmp_path):
    path = book(tmp_path / "p.xlsx", {"A1": 1, "A2": '=XLOOKUP(1,A1:A1,A1:A1)', "A3": '=FILTER(A1:A1,A1:A1>0)'})
    before = path.read_bytes()
    report = run(RECALC, path).json()
    assert report["recalculated"] is False
    fixes = {i["cell"]: i["fix"] for i in report["missing_prefix"]}
    assert fixes == {
        "S!A2": "=_xlfn.XLOOKUP(1,A1:A1,A1:A1)",
        "S!A3": "=_xlfn._xlws.FILTER(A1:A1,A1:A1>0)",
    }
    assert path.read_bytes() == before


def test_prefixed_modern_functions_evaluate(run, tmp_path):
    # TEXTJOIN, because every LibreOffice in use has it; XLOOKUP needs 24.8 or later.
    path = book(tmp_path / "x.xlsx", {"A1": 5, "B1": "five", "A2": '=_xlfn.TEXTJOIN("-",TRUE,A1,B1)'})
    assert run(RECALC, path).code == 0
    assert load_workbook(path, data_only=True)["S"]["A2"].value == "5-five"


def test_circular_references_including_through_names(run, tmp_path):
    path = book(
        tmp_path / "c.xlsx",
        {"A1": "=A2+1", "A2": "=A1+1", "B1": "=loop+1", "B2": "=B1*2", "C1": "=SUM(D1:D9)"},
        names={"loop": "S!$B$2"},
    )
    report = run(RECALC, path).json()
    assert sorted(report["circular_references"]) == [["S!A1", "S!A2"], ["S!B1", "S!B2"]]


def test_csv_is_refused_quickly(run, tmp_path):
    path = tmp_path / "d.csv"
    path.write_text("a,b\n1,2\n")
    result = run(RECALC, path, timeout=30)
    assert result.code == 2 and "xlsx" in result.json()["error"]


def test_external_links_need_force(run, tmp_path):
    path = book(tmp_path / "l.xlsx", {"A1": 1})
    with zipfile.ZipFile(path, "a") as zf:
        zf.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    result = run(RECALC, path)
    assert result.code == 2 and "--force" in result.json()["error"]


def test_users_libreoffice_profile_is_never_touched(run, tmp_path, monkeypatch):
    # Keep uv's cache and Python installs where they are, so only LibreOffice sees the new HOME.
    for var, cmd in (("UV_CACHE_DIR", ["uv", "cache", "dir"]),
                     ("UV_PYTHON_INSTALL_DIR", ["uv", "python", "dir"])):
        monkeypatch.setenv(var, subprocess.check_output(cmd, text=True).strip())
    home = tmp_path / "home"
    home.mkdir()
    monkeypatch.setenv("HOME", str(home))
    path = book(tmp_path / "h.xlsx", {"A1": "=1+1"})
    assert run(RECALC, path).code == 0
    assert not any(home.rglob("*.xba")), "recalc wrote into the user's LibreOffice profile"
