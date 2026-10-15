"""The shared OOXML scripts (unpack, pack, validate, render), as bundled in docx and pptx."""

import re
import zipfile

import pytest
from docx import Document
from pptx import Presentation

OOXML = "docx/ooxml/scripts"
JANE = 'w:author="Jane" w:date="2026-01-01T00:00:00Z"'
BOT = 'w:author="Bot" w:date="2026-02-01T00:00:00Z"'


def docx_with_body(path, body):
    """A .docx whose body is `body` (WordprocessingML paragraphs)."""
    doc = Document()
    doc.add_paragraph("PLACEHOLDER")
    doc.save(path)
    with zipfile.ZipFile(path) as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}
    xml = parts["word/document.xml"].decode()
    parts["word/document.xml"] = re.sub(r"<w:p>(?:(?!<w:p>).)*?PLACEHOLDER.*?</w:p>", body, xml, flags=re.S).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            zf.writestr(name, data)
    return path


def unpack(run, src, dest):
    assert run(f"{OOXML}/unpack.py", src, dest).code == 0
    return dest / "word" / "document.xml"


def edit(path, old, new):
    text = path.read_text(encoding="utf-8")
    assert re.search(old, text, flags=re.S), f"pattern not found: {old}"
    path.write_text(re.sub(old, new, text, count=1, flags=re.S), encoding="utf-8")


def validate(run, target, original, *extra):
    return run(f"{OOXML}/validate.py", target, "--original", original, *extra)


@pytest.fixture
def contract(tmp_path):
    body = '<w:p><w:r><w:t>The term is 30 days.</w:t></w:r></w:p>'
    return docx_with_body(tmp_path / "contract.docx", body)


@pytest.fixture
def reviewed(tmp_path):
    """A document already carrying another author's tracked changes."""
    body = (
        f'<w:p><w:r><w:t xml:space="preserve">Fee is </w:t></w:r>'
        f'<w:ins w:id="1" {JANE}><w:r><w:t>ten</w:t></w:r></w:ins>'
        f'<w:r><w:t xml:space="preserve"> dollars.</w:t></w:r></w:p>'
        f'<w:p><w:del w:id="2" {JANE}><w:r><w:delText xml:space="preserve"> </w:delText></w:r></w:del>'
        f'<w:r><w:t>End.</w:t></w:r></w:p>'
    )
    return docx_with_body(tmp_path / "reviewed.docx", body)


def test_untracked_edit_fails_whatever_the_author(run, tmp_path, contract):
    xml = unpack(run, contract, tmp_path / "u")
    edit(xml, "30 days", "5 days")
    result = validate(run, tmp_path / "u", contract, "--author", "Jane Lawyer")
    assert result.code == 1 and "not tracked" in result.out


def test_tracked_edit_passes(run, tmp_path, contract):
    xml = unpack(run, contract, tmp_path / "u")
    edit(xml, r"<w:r>\s*<w:t>The term is 30 days.</w:t>\s*</w:r>",
         f'<w:r><w:t xml:space="preserve">The term is </w:t></w:r>'
         f'<w:del w:id="9" {BOT}><w:r><w:delText>30</w:delText></w:r></w:del>'
         f'<w:ins w:id="10" {BOT}><w:r><w:t>60</w:t></w:r></w:ins>'
         f'<w:r><w:t xml:space="preserve"> days.</w:t></w:r>')
    result = validate(run, tmp_path / "u", contract)
    assert result.code == 0, result.out


def test_editing_inside_another_authors_insertion_fails(run, tmp_path, reviewed):
    xml = unpack(run, reviewed, tmp_path / "u")
    edit(xml, "<w:t>ten</w:t>", "<w:t>twelve</w:t>")
    result = validate(run, tmp_path / "u", reviewed, "--author", "Bot")
    assert result.code == 1 and "edited directly" in result.out


def test_rejecting_another_authors_insertion_properly_passes(run, tmp_path, reviewed):
    xml = unpack(run, reviewed, tmp_path / "u")
    edit(xml, r"<w:r>\s*<w:t>ten</w:t>\s*</w:r>",
         f'<w:del w:id="11" {BOT}><w:r><w:delText>ten</w:delText></w:r></w:del>')
    edit(xml, r"(</w:ins>)", rf'\1<w:ins w:id="12" {BOT}><w:r><w:t>twelve</w:t></w:r></w:ins>')
    result = validate(run, tmp_path / "u", reviewed, "--author", "Bot")
    assert result.code == 0, result.out


def test_pack_keeps_whitespace_only_text(run, tmp_path, reviewed):
    unpack(run, reviewed, tmp_path / "u")
    out = tmp_path / "repacked.docx"
    assert run(f"{OOXML}/pack.py", tmp_path / "u", out).code == 0
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert '<w:delText xml:space="preserve"> </w:delText>' in xml


def test_literal_non_ascii_text_survives_unpack_edit_pack(run, tmp_path):
    src = docx_with_body(tmp_path / "zh.docx", "<w:p><w:r><w:t>期限为30天</w:t></w:r></w:p>")
    xml = unpack(run, src, tmp_path / "u")
    assert "期限为30天" in xml.read_text(encoding="utf-8")
    edit(xml, r"(<w:t>期限为30天</w:t>\s*</w:r>)", rf'\1<w:ins w:id="5" {BOT}><w:r><w:t>“补充”—条款</w:t></w:r></w:ins>')
    assert validate(run, tmp_path / "u", src).code == 0
    out = tmp_path / "zh-out.docx"
    assert run(f"{OOXML}/pack.py", tmp_path / "u", out).code == 0
    assert "“补充”—条款" in zipfile.ZipFile(out).read("word/document.xml").decode()


@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    slide.placeholders[1].text = "Revenue up 18%"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_validate_checks_slide_xml(run, tmp_path, deck):
    unpack(run, deck, tmp_path / "u")
    assert validate(run, tmp_path / "u", deck).code == 0
    slide = tmp_path / "u" / "ppt" / "slides" / "slide1.xml"
    edit(slide, r"(<p:cSld>)", r"\1<p:bogus/>")
    result = validate(run, tmp_path / "u", deck)
    assert result.code == 1 and "slide1.xml" in result.out


def test_validate_a_generated_file_without_an_original(run, deck):
    result = run(f"{OOXML}/validate.py", deck)
    assert result.code == 0, result.out


@pytest.mark.parametrize("fixture", ["contract", "deck"])
def test_render_produces_page_images(run, tmp_path, fixture, request):
    src = request.getfixturevalue(fixture)
    result = run(f"{OOXML}/render.py", src, tmp_path / "pages")
    images = result.out.split()
    assert result.code == 0 and images and all(i.endswith(".png") for i in images)


def test_shared_copies_are_identical():
    from conftest import ROOT

    def files(root):
        return {p.relative_to(root): p.read_bytes() for p in root.rglob("*")
                if p.is_file() and "__pycache__" not in p.parts}

    assert files(ROOT / "docx/ooxml") == files(ROOT / "pptx/ooxml"), "docx/ooxml and pptx/ooxml differ"
    helper = (ROOT / "docx/ooxml/scripts/soffice.py").read_bytes()
    assert (ROOT / "xlsx/scripts/soffice.py").read_bytes() == helper


def test_inserting_inside_another_authors_insertion_passes(run, tmp_path, reviewed):
    """Word splits their <w:ins> around yours; that is a tracked change, not an edit."""
    xml = unpack(run, reviewed, tmp_path / "u")
    edit(xml, rf'<w:ins w:id="1" {JANE}>\s*<w:r>\s*<w:t>ten</w:t>\s*</w:r>\s*</w:ins>',
         f'<w:ins w:id="1" {JANE}><w:r><w:t>t</w:t></w:r></w:ins>'
         f'<w:ins w:id="20" {BOT}><w:r><w:t>w</w:t></w:r></w:ins>'
         f'<w:ins w:id="21" {JANE}><w:r><w:t>en</w:t></w:r></w:ins>')
    result = validate(run, tmp_path / "u", reviewed, "--author", "Bot")
    assert result.code == 0, result.out


def test_removing_another_authors_insertion_fails(run, tmp_path, reviewed):
    xml = unpack(run, reviewed, tmp_path / "u")
    edit(xml, rf'<w:ins w:id="1" {JANE}>.*?</w:ins>', "")
    result = validate(run, tmp_path / "u", reviewed, "--author", "Bot")
    assert result.code == 1 and "was removed" in result.out


def test_powerpoint_integer_bullet_size_is_valid(run, tmp_path, deck):
    unpack(run, deck, tmp_path / "u")
    slide = tmp_path / "u" / "ppt" / "slides" / "slide1.xml"
    edit(slide, r"<a:p>", '<a:p><a:pPr><a:buSzPct val="100000"/></a:pPr>')
    result = run(f"{OOXML}/validate.py", tmp_path / "u", "--original", deck)
    assert result.code == 0, result.out


def test_absolute_and_external_relationship_targets(run, tmp_path, deck):
    unpack(run, deck, tmp_path / "u")
    rels = tmp_path / "u" / "ppt" / "slides" / "_rels" / "slide1.xml.rels"
    edit(rels, r'Target="\.\./slideLayouts/', 'Target="/ppt/slideLayouts/')
    edit(rels, r"(</Relationships>)",
         r'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
         r'Target="https://example.com/a b" TargetMode="External"/>\1')
    result = run(f"{OOXML}/validate.py", tmp_path / "u", "--original", deck)
    assert result.code == 0, result.out


def test_masters_sharing_a_theme_are_caught(run, tmp_path, deck):
    """PowerPoint repairs a deck whose masters share a theme, though the schema allows it."""
    prs = Presentation(deck)
    prs.slides[0].notes_slide.notes_text_frame.text = "Notes"
    prs.save(deck)
    unpack(run, deck, tmp_path / "u")
    assert run(f"{OOXML}/validate.py", tmp_path / "u").code == 0
    rels = tmp_path / "u" / "ppt" / "notesMasters" / "_rels" / "notesMaster1.xml.rels"
    edit(rels, r"\.\./theme/theme\d+\.xml", "../theme/theme1.xml")
    result = run(f"{OOXML}/validate.py", tmp_path / "u")
    assert result.code == 1 and "share a theme" in result.out
