"""The docx skill: the Document library, and the samples in docx-js.md and ooxml.md.

Library scripts are written to disk and run with `uv run`, the way SKILL.md
tells an agent to run them.
"""

import re
import shutil
import subprocess
import textwrap
import zipfile

import pytest
from docx import Document as PyDocx
from PIL import Image

from conftest import ROOT

SKILL = ROOT / "docx"
OOXML = "docx/ooxml/scripts"
JANE = 'w:author="Jane" w:date="2026-01-01T00:00:00Z"'

needs_node = pytest.mark.skipif(
    shutil.which("node") is None or not (SKILL / "node_modules" / "docx").is_dir(),
    reason="node and docx/node_modules are needed",
)
needs_pandoc = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is needed")


def docx_with_body(path, body):
    """A .docx whose body is `body` (WordprocessingML paragraphs)."""
    doc = PyDocx()
    doc.add_paragraph("PLACEHOLDER")
    doc.save(path)
    with zipfile.ZipFile(path) as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}
    xml = parts["word/document.xml"].decode()
    parts["word/document.xml"] = re.sub(
        r"<w:p>(?:(?!<w:p>).)*?PLACEHOLDER.*?</w:p>", body, xml, flags=re.S
    ).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            zf.writestr(name, data)
    return path


def library(tmp_path, body, deps=("defusedxml>=0.7", "lxml>=5")):
    """Run an agent-style script that imports the Document library."""
    script = tmp_path / "edit.py"
    header = (
        "# /// script\n# requires-python = \">=3.12\"\n"
        f"# dependencies = {list(deps)!r}\n# ///\n".replace("'", '"')
        + f"import sys\nsys.path.insert(0, {str(SKILL)!r})\n"
        "from scripts.document import Document, DocxXMLEditor\n"
    )
    script.write_text(header + textwrap.dedent(body), encoding="utf-8")
    return subprocess.run(
        ["uv", "run", "-q", str(script)], cwd=tmp_path, capture_output=True, text=True, timeout=300
    )


def unpack(run, src, dest="unpacked"):
    assert run(f"{OOXML}/unpack.py", src, dest).code == 0


def pack(run, tmp_path, name="out.docx", src="unpacked"):
    result = run(f"{OOXML}/pack.py", src, name)
    assert result.code == 0, result.err
    return tmp_path / name


def pandoc(path, changes="all"):
    return subprocess.run(
        ["pandoc", str(path), "-t", "markdown", f"--track-changes={changes}"],
        capture_output=True, text=True, check=True,
    ).stdout


def document_xml(path):
    return zipfile.ZipFile(path).read("word/document.xml").decode()


@pytest.fixture
def contract(tmp_path):
    return docx_with_body(tmp_path / "contract.docx", (
        '<w:p><w:r w:rsidR="00112233"><w:rPr><w:rFonts w:eastAsia="SimSun"/><w:b/></w:rPr><w:t>期限为</w:t></w:r>'
        '<w:r w:rsidR="00112233"><w:rPr><w:b/></w:rPr><w:t>30</w:t></w:r>'
        '<w:r><w:rPr><w:rFonts w:eastAsia="SimSun"/><w:b/></w:rPr><w:t>天。</w:t></w:r></w:p>'
        '<w:p><w:r><w:t xml:space="preserve">The </w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>term</w:t></w:r>'
        '<w:r><w:t xml:space="preserve"> is 30 days. The Seller pays.</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Section 5: the Seller ships.</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Delete this paragraph.</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>The end.</w:t></w:r></w:p>'
    ))


@pytest.fixture
def reviewed(tmp_path):
    """Jane's existing changes: an insertion, a deletion, a formatting change,
    a deleted paragraph and an inserted one."""
    return docx_with_body(tmp_path / "reviewed.docx", (
        f'<w:p><w:r><w:t xml:space="preserve">Fee is </w:t></w:r>'
        f'<w:ins w:id="1" {JANE}><w:r><w:t>ten</w:t></w:r></w:ins>'
        f'<w:del w:id="2" {JANE}><w:r><w:delText>five</w:delText></w:r></w:del>'
        f'<w:r><w:rPr><w:b/><w:rPrChange w:id="3" {JANE}><w:rPr><w:i/></w:rPr></w:rPrChange></w:rPr>'
        f'<w:t xml:space="preserve"> dollars.</w:t></w:r></w:p>'
        f'<w:p><w:pPr><w:rPr><w:del w:id="4" {JANE}/></w:rPr></w:pPr>'
        f'<w:del w:id="5" {JANE}><w:r><w:delText>Gone.</w:delText></w:r></w:del></w:p>'
        f'<w:p><w:pPr><w:rPr><w:ins w:id="6" {JANE}/></w:rPr></w:pPr>'
        f'<w:ins w:id="7" {JANE}><w:r><w:t>New para.</w:t></w:r></w:ins></w:p>'
        f'<w:p><w:r><w:t>End.</w:t></w:r></w:p>'
    ))


# ---------- Tracking is enforced on save ----------

def test_untracked_edit_fails_on_save_whatever_the_author(run, tmp_path, contract):
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked", author="Jane Lawyer")
        t = doc["word/document.xml"].get_node(tag="w:t", contains="is 30 days")
        t.firstChild.data = t.firstChild.data.replace("30", "5")
        doc.save()
    """)
    assert result.returncode != 0
    assert "not tracked" in result.stdout and "Redlining validation failed" in result.stderr


def test_require_tracking_false_allows_direct_edits(run, tmp_path, contract):
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked")
        t = doc["word/document.xml"].get_node(tag="w:t", contains="is 30 days")
        t.firstChild.data = t.firstChild.data.replace("30", "5")
        doc.save(require_tracking=False)
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "is 5 days" in (tmp_path / "unpacked/word/document.xml").read_text(encoding="utf-8")


# ---------- tracked_replace ----------

@needs_pandoc
def test_tracked_replace_marks_only_what_differs_across_runs(run, tmp_path, contract):
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked", author="Reviewer")
        body = doc["word/document.xml"]
        body.tracked_replace("期限为30天", "期限为45天")
        body.tracked_replace("term is 30 days", "period is 45 days")
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    out = pack(run, tmp_path)
    xml = document_xml(out)
    deleted = re.findall(r"<w:delText[^>]*>([^<]*)</w:delText>", xml)
    inserted = re.findall(r"<w:ins [^>]*><w:r[^>]*>(?:<w:rPr>.*?</w:rPr>)?<w:t[^>]*>([^<]*)</w:t>", xml)
    assert deleted == ["30", "term", "30"] and inserted == ["45", "period", "45"]
    # The inserted text carries the formatting of the text it replaced
    assert re.search(r"<w:ins [^>]*><w:r[^>]*><w:rPr><w:b/></w:rPr><w:t[^>]*>45<", xml)
    assert re.search(r"<w:ins [^>]*><w:r[^>]*><w:rPr><w:i/></w:rPr><w:t[^>]*>period<", xml)
    accepted = pandoc(out, "accept")
    assert "期限为45天。" in accepted and "The *period* is 45 days." in accepted
    assert "期限为30天。" in pandoc(out, "reject")


def test_tracked_replace_ambiguity_scope_and_replace_all(run, tmp_path, contract):
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked")
        body = doc["word/document.xml"]
        try:
            body.tracked_replace("Seller", "Vendor")
        except ValueError as e:
            print("AMBIGUOUS:", e)
        body.tracked_replace("Seller", "Vendor", scope=body.get_node(tag="w:p", contains="Section 5"))
        print(len(body.tracked_replace("The", "A", replace_all=True)))
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "AMBIGUOUS:" in result.stdout and "2 times" in result.stdout.replace("occurs ", "")


def test_tracked_replace_refuses_to_split_another_authors_insertion(run, tmp_path):
    src = docx_with_body(tmp_path / "jane.docx",
        f'<w:p><w:ins w:id="8" {JANE}><w:r><w:t>in silence, safe and sound</w:t></w:r></w:ins></w:p>')
    unpack(run, src)
    result = library(tmp_path, """
        doc = Document("unpacked")
        body = doc["word/document.xml"]
        try:
            body.tracked_replace("silence", "noise")
        except ValueError as e:
            print("REFUSED:", e)
        body.tracked_replace("safe and sound", "soft and unbound")   # reaches the end: allowed
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "REFUSED:" in result.stdout and "Jane" in result.stdout
    xml = (tmp_path / "unpacked/word/document.xml").read_text(encoding="utf-8")
    assert "".join(re.findall(r"<w:delText[^>]*>([^<]*)<", xml)) == "safe and sound"  # the refusal changed nothing


# ---------- Deleting paragraphs ----------

@needs_pandoc
def test_deleted_paragraph_disappears_on_accept(run, tmp_path, contract):
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked")
        body = doc["word/document.xml"]
        body.suggest_deletion(body.get_node(tag="w:p", contains="Delete this paragraph"))
        body.suggest_deletion(body.get_node(tag="w:p", contains="The end."))   # last: keeps its mark
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    out = pack(run, tmp_path)
    marks = re.findall(r"<w:pPr><w:rPr><w:del ([^>]*)/>", document_xml(out))
    assert len(marks) == 1 and 'w:author="Claude"' in marks[0] and "w:id=" in marks[0]
    accepted = pandoc(out, "accept")
    assert "Delete this paragraph" not in accepted and "Section 5" in accepted


# ---------- Comments ----------

@needs_node
@needs_pandoc
@pytest.mark.parametrize("maker", ["docx-js", "pandoc"])
def test_comments_on_documents_that_already_have_comments_xml(run, tmp_path, maker):
    src = tmp_path / "src.docx"
    if maker == "docx-js":
        (tmp_path / "mk.js").write_text(
            'const fs = require("fs"); const { Document, Packer, Paragraph } = require("docx");\n'
            'Packer.toBuffer(new Document({ sections: [{ children: [new Paragraph("The term is 30 days.")] }] }))'
            '.then(b => fs.writeFileSync("src.docx", b));\n'
        )
        subprocess.run(["node", "mk.js"], cwd=tmp_path, check=True,
                       env={"NODE_PATH": str(SKILL / "node_modules"), "PATH": "/usr/bin:/bin:" + shutil.which("node").rsplit("/", 1)[0]})
    else:
        subprocess.run(["pandoc", "-o", str(src)], input="The term is 30 days.\n", text=True, check=True)
    assert "word/comments.xml" in zipfile.ZipFile(src).namelist()
    unpack(run, src)
    result = library(tmp_path, """
        doc = Document("unpacked")
        p = doc["word/document.xml"].get_node(tag="w:p", contains="30 days")
        cid = doc.add_comment(start=p, end=p, text="Check this")
        doc.reply_to_comment(parent_comment_id=cid, text="Checked")
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    out = pack(run, tmp_path)
    assert run(f"{OOXML}/validate.py", out, "--original", src).code == 0
    text = pandoc(out)
    assert "Check this" in text and "Checked" in text


def test_new_namespaces_are_declared_ignorable(run, tmp_path, contract):
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked")
        doc["word/document.xml"].tracked_replace("The end.", "The very end.")
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    root = re.search(r"<w:document[^>]*>", (tmp_path / "unpacked/word/document.xml").read_text(encoding="utf-8")).group()
    assert "xmlns:w16du=" in root
    assert "w16du" in re.search(r'mc:Ignorable="([^"]*)"', root).group(1).split()


# ---------- Accept / reject all ----------

@pytest.mark.parametrize("mode, expected", [
    ("accept", ["Fee is ten dollars.", "New para.", "End."]),
    ("reject", ["Fee is five dollars.", "Gone.", "End."]),
])
def test_accept_and_reject_all(run, tmp_path, reviewed, mode, expected):
    unpack(run, reviewed)
    result = library(tmp_path, f"""
        doc = Document("unpacked")
        doc.{mode}_all_changes()
        doc.save()
    """)
    assert result.returncode == 0, result.stdout + result.stderr
    out = pack(run, tmp_path)
    xml = document_xml(out)
    assert not re.search(r"<w:(ins|del|rPrChange)\b", xml)
    paragraphs = ["".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", p)) for p in re.findall(r"<w:p\b.*?</w:p>", xml, flags=re.S)]
    assert [p for p in paragraphs if p] == expected
    # Rejecting a formatting change restores the original formatting
    assert ("<w:i/>" in xml) == (mode == "reject")


def test_resolving_changes_refuses_after_untracked_edits(run, tmp_path, reviewed):
    unpack(run, reviewed)
    result = library(tmp_path, """
        doc = Document("unpacked")
        t = doc["word/document.xml"].get_node(tag="w:t", contains="End.")
        t.firstChild.data = "Finish."
        doc.accept_all_changes()
    """)
    assert result.returncode != 0 and "Untracked edits" in result.stderr


# ---------- Imports ----------

def test_import_survives_a_scripts_package_in_the_users_project(run, tmp_path, contract):
    (tmp_path / "scripts").mkdir()
    (tmp_path / "scripts" / "__init__.py").write_text("")
    unpack(run, contract)
    result = library(tmp_path, """
        doc = Document("unpacked")
        print("imported", Document.__module__)
    """)
    assert result.returncode == 0, result.stderr
    assert "imported scripts.document" in result.stdout


# ---------- The samples in the reference docs run as written ----------

def code_blocks(md, lang):
    text = (SKILL / md).read_text(encoding="utf-8")
    return re.findall(rf"^```{lang}\n(.*?)^```", text, flags=re.S | re.M)


@needs_node
def test_docx_js_example_builds_a_valid_letter_document(run, tmp_path):
    Image.new("RGB", (400, 200), "#2E4053").save(tmp_path / "chart.png")
    (tmp_path / "logo.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"><rect width="100" height="50"/></svg>')
    (tmp_path / "build.js").write_text(code_blocks("docx-js.md", "javascript")[0])
    node_dir = shutil.which("node").rsplit("/", 1)[0]
    subprocess.run(["node", "build.js"], cwd=tmp_path, check=True, capture_output=True,
                   env={"NODE_PATH": str(SKILL / "node_modules"), "PATH": f"/usr/bin:/bin:{node_dir}"})
    out = tmp_path / "report.docx"
    result = run(f"{OOXML}/validate.py", out)
    assert result.code == 0, result.out
    xml = document_xml(out)
    assert 'w:w="12240"' in re.search(r"<w:pgSz[^>]*>", xml).group()
    assert '<w:bookmarkStart w:name="results"' in xml.replace('w:id="1" w:name', 'w:name') or 'w:name="results"' in xml
    assert "w:footnoteReference" in xml
    styles = zipfile.ZipFile(out).read("word/styles.xml").decode()
    assert styles.count('w:styleId="Heading1"') == 1
    assert "w:updateFields" in zipfile.ZipFile(out).read("word/settings.xml").decode()


def test_ooxml_image_example_runs_twice_without_duplicates(run, tmp_path, contract):
    sample = next(b for b in code_blocks("ooxml.md", "python") if "pic_id" in b)
    sample = sample.replace("<skill>", str(SKILL)).replace("Second paragraph", "The end.")
    (tmp_path / "img.py").write_text(sample, encoding="utf-8")
    Image.new("RGB", (300, 150), "red").save(tmp_path / "chart.png")
    unpack(run, contract)
    for _ in range(2):
        result = subprocess.run(["uv", "run", "-q", "img.py"], cwd=tmp_path, capture_output=True, text=True, timeout=300)
        assert result.returncode == 0, result.stdout + result.stderr
    out = pack(run, tmp_path)
    assert run(f"{OOXML}/validate.py", out, "--original", contract).code == 0
    parts = zipfile.ZipFile(out)
    assert parts.read("[Content_Types].xml").decode().count('Extension="png"') == 1
    ids = re.findall(r'<wp:docPr id="(\d+)"', document_xml(out))
    assert len(ids) == 2 and len(set(ids)) == 2
