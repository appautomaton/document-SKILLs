"""The pptx skill: html2pptx and the template scripts, run the way SKILL.md says."""

import os
import re
import shutil
import subprocess
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from conftest import ROOT

SKILL = ROOT / "pptx"
VALIDATE = "pptx/ooxml/scripts/validate.py"
needs_node = pytest.mark.skipif(
    shutil.which("node") is None or not (SKILL / "node_modules").is_dir(),
    reason="needs node and `npm install` in pptx/",
)

SLIDE_CSS = """
body { width: 720pt; height: 405pt; margin: 0; display: flex; flex-direction: column;
       font-family: Arial; background: #FAFAF7; }
h1 { font-size: 30pt; margin: 36pt 48pt 12pt; }
p, ul, ol { font-size: 16pt; margin: 0 48pt 8pt; }
"""


def slide_html(body, css=SLIDE_CSS):
    return f"<!DOCTYPE html><html><head><style>{css}</style></head><body>{body}</body></html>"


def node(script, cwd):
    """Run a build script from the user's directory, with the skill's packages on NODE_PATH."""
    env = {**os.environ, "NODE_PATH": str(SKILL / "node_modules")}
    return subprocess.run(["node", script], cwd=cwd, env=env, capture_output=True, text=True, timeout=300)


def build_script(body):
    return (
        "const pptxgen = require('pptxgenjs');\n"
        f"const html2pptx = require({str(SKILL / 'scripts' / 'html2pptx.js')!r});\n"
        "async function main() {\n"
        "  const pptx = new pptxgen();\n  pptx.layout = 'LAYOUT_16x9';\n"
        f"{body}\n"
        "}\n"
        "main().catch((e) => { console.error(e.message); process.exit(1); });\n"
    )


def part(path, name):
    return zipfile.ZipFile(path).read(name).decode()


@pytest.fixture
def workdir(tmp_path):
    """A working directory whose path contains a space, as users' often do."""
    path = tmp_path / "my deck"
    path.mkdir()
    return path


@needs_node
def test_html2pptx_deck_is_valid_and_keeps_formatting(run, workdir):
    Image.new("RGB", (200, 120), "#C0392B").save(workdir / "red box.png")
    (workdir / "s1.html").write_text(slide_html(
        '<h1>Quarterly <span style="color:#C0392B">review</span></h1>'
        "<p>Revenue up <b>18%</b>, margin <i>flat</i>, <u>churn</u> down.<br>Second line.</p>"
        "<ul><li>First <b>bold</b> point</li><li>Second point<ul><li>Nested detail</li></ul></li></ul>"
        "<ol><li>Step one</li><li>Step <b>two</b></li><li>Step three</li></ol>"
    ))
    (workdir / "s2.html").write_text(slide_html(
        '<div style="display:flex;margin:40pt">'
        '<div style="width:40%"><p style="margin:0">Chart on the right.</p>'
        '<img src="red box.png" style="width:120pt;height:72pt;margin-top:16pt">'
        '<div style="background:#F2EEE6;border-left:6pt solid #C0392B;padding:12pt;margin-top:16pt">'
        '<p style="margin:0">Callout</p></div></div>'
        '<div style="width:60%"><div id="chart" class="placeholder" style="width:100%;height:220pt"></div></div>'
        "</div>"
    ))
    (workdir / "build.js").write_text(build_script("""
  await html2pptx('s1.html', pptx);
  const { slide, placeholders } = await html2pptx('s2.html', pptx);
  slide.addChart(pptx.charts.BAR, [{ name: 'Revenue', labels: ['Q1','Q2'], values: [4.5, 5.5] }],
    { ...placeholders[0], chartColors: ['C0392B'] });
  slide.addNotes('Speaker notes');
  await pptx.writeFile({ fileName: 'deck.pptx' });"""))

    built = node("build.js", workdir)
    assert built.returncode == 0, built.stderr
    deck = workdir / "deck.pptx"
    result = run(VALIDATE, deck, cwd=workdir)
    assert result.code == 0, result.out

    slide1 = part(deck, "ppt/slides/slide1.xml")
    for paragraph in re.findall(r"<a:p>.*?</a:p>", slide1, flags=re.S):
        assert paragraph.count("<a:pPr") <= 1
    assert [int(n) for n in re.findall(r'<a:buAutoNum type="arabicPeriod" startAt="(\d+)"', slide1)] == [1, 2, 3]
    assert "<a:br/>" in slide1 and 'u="sng"' in slide1 and 'lvl="1"' in slide1
    # Numbers are wider than bullets, so they get more room before the text.
    hanging = lambda marker: int(re.search(rf'marL="(\d+)"[^>]*>(?:(?!</a:pPr>).)*{marker}', slide1, re.S).group(1))
    assert hanging("<a:buAutoNum") > hanging("<a:buChar")
    assert any(n.startswith("ppt/media/") for n in zipfile.ZipFile(deck).namelist())

    pages = run("pptx/ooxml/scripts/render.py", deck, workdir / "pages", cwd=workdir)
    assert pages.code == 0 and len(pages.out.splitlines()) == 2


@needs_node
def test_one_browser_serves_every_slide(workdir):
    for i in range(3):
        (workdir / f"s{i}.html").write_text(slide_html(f"<h1>Slide {i}</h1>"))
    (workdir / "build.js").write_text(
        "const { chromium } = require('playwright-core');\n"
        "let launches = 0;\n"
        "const launch = chromium.launch.bind(chromium);\n"
        "chromium.launch = async (o) => { const b = await launch(o); launches++; return b; };\n"
        + build_script("""
  for (let i = 0; i < 3; i++) await html2pptx(`s${i}.html`, pptx);
  console.log(launches);""")
    )
    built = node("build.js", workdir)
    assert built.returncode == 0, built.stderr
    assert built.stdout.strip() == "1"


@needs_node
def test_pptxgenjs_charts_and_mixed_runs_come_out_valid(run, workdir):
    (workdir / "s.html").write_text(slide_html(
        '<div id="chart" class="placeholder" style="width:600pt;height:250pt;margin:40pt"></div>'
    ))
    (workdir / "build.js").write_text(build_script("""
  const data = [{ name: 'A', labels: ['Q1','Q2','Q3'], values: [1,2,3] },
                { name: 'B', labels: ['Q1','Q2','Q3'], values: [2,1,4] }];
  const xy = [{ name: 'X', values: [1,2,3] }, { name: 'Y', values: [3,1,2] }];
  for (const type of ['LINE', 'AREA', 'SCATTER', 'RADAR', 'BAR', 'PIE']) {
    const { slide, placeholders } = await html2pptx('s.html', pptx);
    const series = type === 'SCATTER' ? xy : type === 'PIE' ? [data[0]] : data;
    slide.addChart(pptx.charts[type], series, { ...placeholders[0], showValue: true,
      barGrouping: 'stacked', dataLabelPosition: type === 'BAR' ? 'outEnd' : undefined,
      showCatAxisTitle: true, catAxisTitle: 'Quarter' });
  }
  // One highlighted bar, as design.md recommends: a colour per data point.
  const { slide, placeholders } = await html2pptx('s.html', pptx);
  slide.addChart(pptx.charts.BAR, [data[0]], { ...placeholders[0], showValue: true,
    chartColors: ['B0B0B0', 'C0392B', 'B0B0B0'] });
  pptx.addSlide().addText([{ text: 'Bold ', options: { bold: true } }, { text: 'plain' }],
    { x: 1, y: 1, w: 6, h: 1 });
  await pptx.writeFile({ fileName: 'charts.pptx' });"""))
    built = node("build.js", workdir)
    assert built.returncode == 0, built.stderr
    result = run(VALIDATE, workdir / "charts.pptx", cwd=workdir)
    assert result.code == 0, result.out


@needs_node
def test_slide_problems_are_reported_together(workdir):
    (workdir / "bad.html").write_text(slide_html(
        "<div>bare text</div><p>• typed bullet</p>",
        css=SLIDE_CSS + "body { background: linear-gradient(#fff, #000); }",
    ))
    (workdir / "build.js").write_text(build_script("  await html2pptx('bad.html', pptx);"))
    built = node("build.js", workdir)
    assert built.returncode == 1
    for fragment in ("bad.html: Multiple validation errors", "unwrapped text", "gradients", "bullet symbol"):
        assert fragment in built.stderr


@needs_node
def test_placeholders_can_be_captured_after_scripts_draw(workdir):
    (workdir / "s.html").write_text(slide_html(
        '<div id="viz" class="placeholder" style="width:300pt;height:200pt;margin:40pt">'
        '<svg width="100%" height="100%"><rect id="r" width="100%" height="100%" fill="#1F3A5F"/></svg></div>'
        "<script>setTimeout(() => document.getElementById('r').setAttribute('fill', '#C0392B'), 100)</script>"
    ))
    (workdir / "build.js").write_text(build_script("""
  const { placeholders } = await html2pptx('s.html', pptx, { screenshotPlaceholders: true, renderDelay: 400 });
  require('fs').writeFileSync('viz.png', Buffer.from(placeholders[0].image.split(',')[1], 'base64'));"""))
    built = node("build.js", workdir)
    assert built.returncode == 0, built.stderr
    with Image.open(workdir / "viz.png") as image:
        r, g, b = image.convert("RGB").getpixel((image.width // 2, image.height // 2))
    assert r > 150 and b < 100, "captured before the script redrew the placeholder"


@pytest.fixture
def chart_template(tmp_path):
    """A template whose first slide has a chart, a hyperlink, a picture and notes."""
    Image.new("RGB", (60, 40), "navy").save(tmp_path / "logo.png")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Chart slide"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (1, 2))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data)
    run_ = slide.shapes.add_textbox(Inches(6.5), Inches(2), Inches(3), Inches(1)).text_frame.paragraphs[0].add_run()
    run_.text = "link"
    run_.hyperlink.address = "https://example.com"
    slide.shapes.add_picture(str(tmp_path / "logo.png"), Inches(7), Inches(4))
    slide.notes_slide.notes_text_frame.text = "Chart notes"
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Second"
    path = tmp_path / "template.pptx"
    prs.save(path)
    # python-pptx writes negative chart axis ids, which the schema rejects; start from a valid template.
    with zipfile.ZipFile(path) as zf:
        parts = {n: zf.read(n) for n in zf.namelist()}
    for name in parts:
        if name.startswith("ppt/charts/chart"):
            parts[name] = re.sub(rb'val="-(\d+)"', rb'val="\1"', parts[name])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in parts.items():
            zf.writestr(name, blob)
    return path


def test_rearrange_duplicates_charts_links_and_notes(run, tmp_path, chart_template):
    out = tmp_path / "out.pptx"
    assert run("pptx/scripts/rearrange.py", chart_template, out, "0,0,1").code == 0
    result = run(VALIDATE, out, "--original", chart_template)
    assert result.code == 0, result.out

    names = zipfile.ZipFile(out).namelist()
    assert len([n for n in names if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]) == 2
    assert len([n for n in names if n.startswith("ppt/embeddings/")]) == 2
    first, duplicate, _ = Presentation(out).slides
    kinds = sorted(r.reltype.rsplit("/", 1)[-1] for r in duplicate.part.rels.values())
    assert kinds == ["chart", "hyperlink", "image", "notesSlide", "slideLayout"]
    assert duplicate.notes_slide.notes_text_frame.text == "Chart notes"
    assert duplicate.notes_slide.part is not first.notes_slide.part


@pytest.fixture
def text_deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Old title"
    body = slide.placeholders[1].text_frame
    body.text = "Old body"
    body.paragraphs[0].line_spacing = Pt(20)
    body.paragraphs[0].space_before = Pt(6)
    keep = prs.slides.add_slide(prs.slide_layouts[1])
    keep.shapes.title.text = "Keep me"
    keep.placeholders[1].text = "Keep body"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_replace_is_schema_valid_and_leaves_unlisted_slides_alone(run, tmp_path, text_deck):
    (tmp_path / "rep.json").write_text("""{"slide-0": {
        "shape-0": {"paragraphs": [{"text": "New title", "bold": true}]},
        "shape-1": {"paragraphs": [{"text": "Plain", "space_before": 6, "line_spacing": 20},
                                   {"text": "A bullet", "bullet": true, "level": 0}]}}}""")
    out = tmp_path / "out.pptx"
    assert run("pptx/scripts/replace.py", text_deck, tmp_path / "rep.json", out).code == 0
    result = run(VALIDATE, out, "--original", text_deck)
    assert result.code == 0, result.out
    texts = [[s.text_frame.text for s in slide.shapes] for slide in Presentation(out).slides]
    assert texts == [["New title", "Plain\nA bullet"], ["Keep me", "Keep body"]]


def test_replace_names_the_shapes_that_exist(run, tmp_path, text_deck):
    (tmp_path / "bad.json").write_text('{"slide-0": {"shape-9": {"paragraphs": [{"text": "x"}]}}}')
    result = run("pptx/scripts/replace.py", text_deck, tmp_path / "bad.json", tmp_path / "x.pptx")
    assert result.code == 1 and "Traceback" not in result.out + result.err
    assert "Available shapes: shape-0 ('Old title'), shape-1 ('Old body')" in result.out


def test_inventory_places_grouped_shapes_where_they_are(run, tmp_path):
    prs = Presentation()
    group = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_group_shape()
    group.shapes.add_textbox(Inches(8), Inches(5), Inches(1.5), Inches(0.5)).text_frame.text = "Grouped"
    prs.save(tmp_path / "g.pptx")
    assert run("pptx/scripts/inventory.py", tmp_path / "g.pptx", tmp_path / "inv.json").code == 0
    shape = __import__("json").loads((tmp_path / "inv.json").read_text())["slide-0"]["shape-0"]
    assert (shape["left"], shape["top"], shape["width"]) == (8.0, 5.0, 1.5)
    assert "overflow" not in shape


def test_thumbnail_rejects_a_bad_column_count(run, text_deck):
    result = run("pptx/scripts/thumbnail.py", text_deck, "--cols", "0")
    assert result.code == 2 and "invalid choice" in result.err


def test_notes_round_trip(run, tmp_path, text_deck):
    (tmp_path / "notes.json").write_text('{"slide-1": "第一行\\nSecond line"}', encoding="utf-8")
    out = tmp_path / "noted.pptx"
    assert run("pptx/scripts/notes.py", text_deck, tmp_path / "notes.json", "-o", out).code == 0
    assert run("pptx/scripts/notes.py", out).json() == {"slide-0": "", "slide-1": "第一行\nSecond line"}
    assert run(VALIDATE, out, "--original", text_deck).code == 0

    (tmp_path / "bad.json").write_text('{"slide-7": "x"}')
    result = run("pptx/scripts/notes.py", text_deck, tmp_path / "bad.json")
    assert result.code == 1 and "slide-0 to slide-1" in result.err


def test_missing_doctype_is_named_instead_of_misreported(workdir):
    (workdir / "q.html").write_text(
        "<html><body style='width:720pt;height:405pt;margin:0'><h1 style='margin:40pt'>Hi</h1></body></html>")
    (workdir / "build.js").write_text(build_script("  await html2pptx('q.html', pptx);"))
    result = node(workdir / "build.js", workdir)
    assert result.returncode == 1
    assert "<!DOCTYPE html>" in result.stderr and "overflows" not in result.stderr
